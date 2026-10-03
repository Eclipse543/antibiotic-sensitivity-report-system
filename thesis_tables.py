import os
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from excel_export import load_ast_data
from resistance_analysis import (
    create_mdr_dataset,
    add_mdr_classification
)


# ---------------------------------------------------------
# Paths
# ---------------------------------------------------------

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "reports"
)

os.makedirs(
    OUTPUT_DIR,
    exist_ok=True
)

OUTPUT_FILE = os.path.join(
    OUTPUT_DIR,
    "Thesis_Statistical_Tables.xlsx"
)


# ---------------------------------------------------------
# Percentage
# ---------------------------------------------------------

def pct(n, total):

    if total == 0:
        return 0.0

    return round(
        n / total * 100,
        2
    )


# ---------------------------------------------------------
# Table 1
# Organism distribution
# ---------------------------------------------------------

def table_1_organism_distribution():

    data = load_ast_data()

    if data.empty:
        return pd.DataFrame()

    isolates = (
        data[
            [
                "isolate_id",
                "organism"
            ]
        ]
        .drop_duplicates()
    )

    total = len(isolates)

    result = (
        isolates["organism"]
        .value_counts()
        .reset_index()
    )

    result.columns = [
        "Organism",
        "Frequency"
    ]

    result["Percentage"] = (
        result["Frequency"]
        .apply(
            lambda x: pct(x, total)
        )
    )

    return result


# ---------------------------------------------------------
# Table 2
# Specimen distribution
# ---------------------------------------------------------

def table_2_specimen_distribution():

    data = load_ast_data()

    if data.empty:
        return pd.DataFrame()

    isolates = (
        data[
            [
                "isolate_id",
                "specimen"
            ]
        ]
        .drop_duplicates()
    )

    total = len(isolates)

    result = (
        isolates["specimen"]
        .value_counts()
        .reset_index()
    )

    result.columns = [
        "Specimen",
        "Frequency"
    ]

    result["Percentage"] = (
        result["Frequency"]
        .apply(
            lambda x: pct(x, total)
        )
    )

    return result


# ---------------------------------------------------------
# Table 3
# Organism by specimen
# ---------------------------------------------------------

def table_3_organism_by_specimen():

    data = load_ast_data()

    if data.empty:
        return pd.DataFrame()

    isolates = (
        data[
            [
                "isolate_id",
                "organism",
                "specimen"
            ]
        ]
        .drop_duplicates()
    )

    table = pd.crosstab(
        isolates["specimen"],
        isolates["organism"]
    )

    return table.reset_index()


# ---------------------------------------------------------
# Table 4
# MDR distribution
# ---------------------------------------------------------

def table_4_mdr_distribution():

    # IMPORTANT:
    # create_mdr_dataset() calculates
    # resistant_class_count before
    # MDR classification.

    data = create_mdr_dataset()

    if data.empty:
        return pd.DataFrame()

    mdr = add_mdr_classification(
        data
    )

    isolates = (
        mdr[
            [
                "isolate_id",
                "organism",
                "MDR_status"
            ]
        ]
        .drop_duplicates()
    )

    result = (
        isolates["MDR_status"]
        .value_counts()
        .reset_index()
    )

    result.columns = [
        "MDR Status",
        "Isolates"
    ]

    total = result["Isolates"].sum()

    result["Percentage"] = (
        result["Isolates"] / total * 100
    ).round(2)

    return result


# ---------------------------------------------------------
# Table 5
# MDR by organism
# ---------------------------------------------------------

def table_5_mdr_by_organism():

    data = create_mdr_dataset()

    if data.empty:
        return pd.DataFrame()

    mdr = add_mdr_classification(
        data
    )

    isolates = (
        mdr[
            [
                "isolate_id",
                "organism",
                "MDR_status"
            ]
        ]
        .drop_duplicates()
    )

    rows = []

    for organism, group in isolates.groupby(
        "organism"
    ):

        total = len(group)

        mdr_count = (
            group["MDR_status"]
            .eq("MDR")
            .sum()
        )

        rows.append({
            "Organism": organism,
            "Total": total,
            "MDR": mdr_count,
            "MDR %": pct(
                mdr_count,
                total
            )
        })

    return pd.DataFrame(rows)


# ---------------------------------------------------------
# Format workbook
# ---------------------------------------------------------

def format_workbook():

    workbook = load_workbook(
        OUTPUT_FILE
    )

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"

        # Header formatting
        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Center numeric cells
        for row in worksheet.iter_rows(
            min_row=2
        ):

            for cell in row:

                if isinstance(
                    cell.value,
                    (int, float)
                ):

                    cell.alignment = Alignment(
                        horizontal="center"
                    )

        # Automatic column width
        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                40
            )

    workbook.save(
        OUTPUT_FILE
    )


# ---------------------------------------------------------
# Generate thesis workbook
# ---------------------------------------------------------

def generate_thesis_tables():

    tables = {

        "Table_1_Organism": (
            table_1_organism_distribution()
        ),

        "Table_2_Specimen": (
            table_2_specimen_distribution()
        ),

        "Table_3_Organism_Specimen": (
            table_3_organism_by_specimen()
        ),

        "Table_4_MDR": (
            table_4_mdr_distribution()
        ),

        "Table_5_MDR_Organism": (
            table_5_mdr_by_organism()
        ),
    }

    # Remove old workbook before creating
    # the new workbook so obsolete sheets
    # cannot remain.

    if os.path.exists(
        OUTPUT_FILE
    ):

        os.remove(
            OUTPUT_FILE
        )

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        for sheet_name, table in tables.items():

            if (
                table is not None
                and not table.empty
            ):

                table.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False
                )

    format_workbook()

    return OUTPUT_FILE


# ---------------------------------------------------------
# Run
# ---------------------------------------------------------

if __name__ == "__main__":

    output = generate_thesis_tables()

    print(
        "\nThesis tables generated successfully:"
    )

    print(output)
